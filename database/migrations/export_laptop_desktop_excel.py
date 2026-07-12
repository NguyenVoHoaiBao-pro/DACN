#!/usr/bin/env python3
"""Xuat 673 serial Laptop & Desktop ra file Excel."""

import ssl
from pathlib import Path

import pymysql
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[2]
ENV_FILE = REPO_ROOT / ".env"
OUT_FILE = REPO_ROOT / "database" / "exports" / "laptop_desktop_serials.xlsx"
CATEGORY_ID = 12


def load_env() -> dict[str, str]:
    env: dict[str, str] = {}
    for raw in ENV_FILE.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        env[key.strip()] = value.strip()
    return env


def main() -> int:
    env = load_env()
    ssl_ctx = ssl.create_default_context()
    ssl_ctx.check_hostname = False
    ssl_ctx.verify_mode = ssl.CERT_NONE
    conn = pymysql.connect(
        host=env["MYSQL_HOST"],
        port=int(env["MYSQL_PORT"]),
        user=env["MYSQL_USER"],
        password=env["MYSQL_PASSWORD"],
        database="electro_catalog_db",
        charset="utf8mb4",
        ssl=ssl_ctx,
    )

    sql = """
        SELECT
            ROW_NUMBER() OVER (ORDER BY p.name, v.sku_code, pi.created_at) AS stt,
            pt.name AS danh_muc,
            p.name AS ten_san_pham,
            v.sku_code,
            v.variant_name,
            v.stock_quantity,
            pi.serial_number,
            pi.imei,
            pi.location,
            pi.batch_number,
            pi.status
        FROM product_items pi
        JOIN product_variants v ON v.id = pi.variant_id
        JOIN products p ON p.id = v.product_id
        JOIN product_types pt ON pt.id = p.product_type_id
        WHERE pi.status = 'AVAILABLE'
          AND pi.serial_number IS NOT NULL
          AND (pt.id = %s OR pt.parent_id = %s)
        ORDER BY p.name, v.sku_code, pi.created_at ASC
    """

    try:
        with conn.cursor() as cur:
            cur.execute(sql, (CATEGORY_ID, CATEGORY_ID))
            rows = cur.fetchall()
    finally:
        conn.close()

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # --- Sheet chi tiet ---
    ws = wb.active
    ws.title = "Danh sach Serial"
    headers = [
        "STT",
        "Danh muc",
        "Ten san pham",
        "SKU",
        "Bien the",
        "Ton kho web",
        "Serial",
        "IMEI",
        "Vi tri ke",
        "Lo hang",
        "Trang thai",
    ]
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, row in enumerate(rows, 2):
        for c_idx, val in enumerate(row, 1):
            ws.cell(row=r_idx, column=c_idx, value=val)

    widths = [6, 18, 42, 22, 12, 12, 22, 22, 20, 16, 12]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    # --- Sheet tom tat theo SKU ---
    ws2 = wb.create_sheet("Tom tat theo SKU")
    ws2.append(["SKU", "Ten san pham", "Bien the", "Ton kho web", "So serial"])
    for col in range(1, 6):
        c = ws2.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font

    summary: dict[str, list] = {}
    for row in rows:
        sku = row[3]
        summary.setdefault(sku, []).append(row)

    for sku in sorted(summary.keys()):
        items = summary[sku]
        first = items[0]
        ws2.append([sku, first[2], first[4], first[5], len(items)])

    for i, w in enumerate([22, 42, 12, 12, 12], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(OUT_FILE)
    print(f"Exported {len(rows)} serials -> {OUT_FILE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
